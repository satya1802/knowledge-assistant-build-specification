"""Tests for app.services.ingestion: extraction and chunking (US-010-1)."""

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.services.ingestion import extract as extract_module
from app.services.ingestion.chunking import chunk_pages
from app.services.ingestion.extract import ExtractedPage, ExtractionError, extract_document


def _write_text_pdf(path: Path, page_texts: list[str]) -> None:
    """Build a minimal, hand-assembled text-based PDF with one page per
    string in `page_texts`. Written as raw PDF syntax (rather than through a
    higher-level writer API) so the test has full control over producing
    real, extractable text content streams with a correct xref table."""
    objects: list[bytes] = []

    def add_object(body: bytes) -> int:
        objects.append(body)
        return len(objects)  # 1-based object number

    font_obj_num = add_object(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    page_obj_nums: list[int] = []
    content_obj_nums: list[int] = []
    pages_obj_num = len(page_texts) * 2 + 3  # reserved, filled in below

    # Reserve page and content object numbers up front so cross-references
    # (Parent, Contents, Font) can be written correctly on the first pass.
    next_num = font_obj_num + 1
    for _ in page_texts:
        page_obj_nums.append(next_num)
        next_num += 1
        content_obj_nums.append(next_num)
        next_num += 1
    pages_obj_num = next_num
    catalog_obj_num = pages_obj_num + 1

    for text, page_num, content_num in zip(
        page_texts, page_obj_nums, content_obj_nums, strict=True
    ):
        objects.append(
            (
                f"<< /Type /Page /Parent {pages_obj_num} 0 R "
                f"/Resources << /Font << /F1 {font_obj_num} 0 R >> >> "
                f"/MediaBox [0 0 200 200] /Contents {content_num} 0 R >>"
            ).encode()
        )
        assert len(objects) == page_num

        stream_body = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode()
        objects.append(
            f"<< /Length {len(stream_body)} >>\nstream\n".encode() + stream_body + b"\nendstream"
        )
        assert len(objects) == content_num

    kids = " ".join(f"{n} 0 R" for n in page_obj_nums)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_obj_nums)} >>".encode())
    assert len(objects) == pages_obj_num

    objects.append(f"<< /Type /Catalog /Pages {pages_obj_num} 0 R >>".encode())
    assert len(objects) == catalog_obj_num

    buf = bytearray()
    buf += b"%PDF-1.4\n"
    offsets = [0]
    for i, body in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_offset = len(buf)
    buf += f"xref\n0 {len(objects) + 1}\n".encode()
    buf += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        buf += f"{off:010d} 00000 n \n".encode()

    buf += (
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_obj_num} 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode()

    path.write_bytes(bytes(buf))


@pytest.fixture
def pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "doc.pdf"
    _write_text_pdf(path, ["Hello page one", "Hello page two"])
    return path


def test_extract_pdf_preserves_page_numbers(pdf_path: Path) -> None:
    pages = extract_document(pdf_path, "pdf")
    assert len(pages) == 2
    assert [p.page_number for p in pages] == [1, 2]
    for page in pages:
        assert isinstance(page, ExtractedPage)


def test_extract_pdf_corrupt_raises_extraction_error(tmp_path: Path) -> None:
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"%PDF-1.4 not actually a pdf")
    with pytest.raises(ExtractionError) as exc_info:
        extract_document(bad, "pdf")
    assert exc_info.value.reason


@pytest.fixture
def scanned_pdf_path(tmp_path: Path) -> Path:
    """A valid PDF with one blank page: pypdf extracts no text from it at
    all, exactly the "almost no text" case a scanned page produces (AC-035)."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "scanned.pdf"
    with open(path, "wb") as fh:
        writer.write(fh)
    return path


def test_extract_pdf_scanned_page_uses_ocr_when_available(
    monkeypatch: pytest.MonkeyPatch, scanned_pdf_path: Path
) -> None:
    monkeypatch.setattr(extract_module, "_ocr_runtime_available", lambda: True)
    monkeypatch.setattr(
        extract_module, "_ocr_page_text", lambda path, page_number: "Recognised scan text"
    )

    pages = extract_document(scanned_pdf_path, "pdf")

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].text == "Recognised scan text"
    assert pages[0].ocr_unavailable is False

    # The recognised text flows into the same chunking path as any other
    # page (AC-035/AC-036): document id is attached by the pipeline, the
    # 1-based page number already survives chunking here.
    chunks = chunk_pages(pages)
    assert len(chunks) == 1
    assert chunks[0].page_number == 1
    assert chunks[0].text == "Recognised scan text"


def test_extract_pdf_scanned_page_skipped_when_ocr_unavailable(
    monkeypatch: pytest.MonkeyPatch, scanned_pdf_path: Path
) -> None:
    monkeypatch.setattr(extract_module, "_ocr_runtime_available", lambda: False)

    pages = extract_document(scanned_pdf_path, "pdf")

    assert len(pages) == 1
    assert pages[0].page_number == 1
    assert pages[0].text == ""
    assert pages[0].ocr_unavailable is True

    # A fully-blank skipped page contributes no chunk, never an error.
    assert chunk_pages(pages) == []


def test_extract_pdf_ocr_failure_degrades_to_skipped_not_a_crash(
    monkeypatch: pytest.MonkeyPatch, scanned_pdf_path: Path
) -> None:
    """`_ocr_page_text` itself swallows render/recognition errors and
    returns `None` (AC-037) -- the page is then marked skipped, never an
    unhandled exception bubbling out of `extract_document`."""
    monkeypatch.setattr(extract_module, "_ocr_runtime_available", lambda: True)
    monkeypatch.setattr(extract_module, "_ocr_page_text", lambda path, page_number: None)

    pages = extract_document(scanned_pdf_path, "pdf")

    assert pages[0].ocr_unavailable is True
    assert pages[0].text == ""


def test_ocr_page_text_returns_none_when_renderer_raises(
    monkeypatch: pytest.MonkeyPatch, scanned_pdf_path: Path
) -> None:
    """Exercises `_ocr_page_text`'s own error handling directly, with fake
    renderer/OCR modules standing in for pypdfium2/pytesseract."""

    class _ExplodingPdfium:
        @staticmethod
        def PdfDocument(path: str):  # noqa: N802 -- matches pypdfium2's API name
            raise RuntimeError("renderer exploded")

    monkeypatch.setattr(extract_module, "_pdfium", _ExplodingPdfium())

    result = extract_module._ocr_page_text(scanned_pdf_path, 1)

    assert result is None


def test_ocr_runtime_available_false_when_modules_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract_module, "_pdfium", None)
    monkeypatch.setattr(extract_module, "_pytesseract", None)
    assert extract_module._ocr_runtime_available() is False


def test_ocr_runtime_available_false_when_tesseract_binary_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class _FakePdfium:
        pass

    class _FakeTesseractMissing:
        @staticmethod
        def get_tesseract_version() -> str:
            raise FileNotFoundError("tesseract is not installed")

    monkeypatch.setattr(extract_module, "_pdfium", _FakePdfium())
    monkeypatch.setattr(extract_module, "_pytesseract", _FakeTesseractMissing())
    assert extract_module._ocr_runtime_available() is False


def test_extract_pdf_password_protected_raises_extraction_error(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="secret", owner_password="secret-owner")
    path = tmp_path / "locked.pdf"
    with open(path, "wb") as fh:
        writer.write(fh)

    with pytest.raises(ExtractionError) as exc_info:
        extract_document(path, "pdf")
    assert "password" in exc_info.value.reason.lower()


def test_extract_docx_paragraphs_and_table_cells_in_order(tmp_path: Path) -> None:
    doc = DocxDocument()
    doc.add_paragraph("First paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell-only-content"
    table.rows[0].cells[1].text = "cell-two"
    doc.add_paragraph("Last paragraph")
    path = tmp_path / "doc.docx"
    doc.save(path)

    pages = extract_document(path, "docx")
    assert len(pages) == 1
    text = pages[0].text
    assert pages[0].page_number is None
    # Content that exists only in a table cell must be present and
    # attributable (AC-033).
    assert "cell-only-content" in text
    assert "cell-two" in text
    assert text.index("First paragraph") < text.index("cell-only-content")
    assert text.index("cell-only-content") < text.index("Last paragraph")


def test_extract_docx_corrupt_raises_extraction_error(tmp_path: Path) -> None:
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a real docx file")
    with pytest.raises(ExtractionError):
        extract_document(bad, "docx")


def test_extract_txt_reads_utf8(tmp_path: Path) -> None:
    path = tmp_path / "doc.txt"
    path.write_text("hello world\n\nmore text", encoding="utf-8")
    pages = extract_document(path, "txt")
    assert len(pages) == 1
    assert pages[0].page_number is None
    assert "hello world" in pages[0].text


def test_extract_md_reads_utf8(tmp_path: Path) -> None:
    path = tmp_path / "doc.md"
    path.write_text("# Heading\n\nSome content", encoding="utf-8")
    pages = extract_document(path, "md")
    assert len(pages) == 1
    assert "Heading" in pages[0].text


def test_extract_txt_empty_line_only_file_does_not_error(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("\n\n\n", encoding="utf-8")
    pages = extract_document(path, "txt")
    assert len(pages) == 1
    # Chunking such a page must not error either.
    chunks = chunk_pages(pages)
    assert chunks == []


def test_extract_txt_very_short_file(tmp_path: Path) -> None:
    path = tmp_path / "short.txt"
    path.write_text("hi", encoding="utf-8")
    pages = extract_document(path, "txt")
    chunks = chunk_pages(pages)
    assert len(chunks) == 1
    assert chunks[0].text == "hi"
    assert chunks[0].page_number is None
    assert chunks[0].ordinal == 0


def test_extract_txt_undecodable_raises_extraction_error(tmp_path: Path) -> None:
    path = tmp_path / "bad.txt"
    path.write_bytes(b"\xff\xfe\xfa\x00invalid-utf8")
    with pytest.raises(ExtractionError) as exc_info:
        extract_document(path, "txt")
    assert exc_info.value.reason


def test_extract_unsupported_file_type_raises_extraction_error(tmp_path: Path) -> None:
    path = tmp_path / "doc.xyz"
    path.write_text("content", encoding="utf-8")
    with pytest.raises(ExtractionError):
        extract_document(path, "xyz")


def test_chunk_pages_assigns_sequential_ordinals_across_pages() -> None:
    pages = [
        ExtractedPage(page_number=1, text="Paragraph A."),
        ExtractedPage(page_number=2, text="Paragraph B."),
    ]
    chunks = chunk_pages(pages)
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))
    assert chunks[0].page_number == 1
    assert chunks[-1].page_number == 2


def test_chunk_pages_is_deterministic() -> None:
    pages = [ExtractedPage(page_number=1, text="word " * 500)]
    first = chunk_pages(pages)
    second = chunk_pages(pages)
    assert [c.text for c in first] == [c.text for c in second]
    assert len(first) > 1


def test_chunk_pages_empty_input_returns_empty_list() -> None:
    assert chunk_pages([]) == []
